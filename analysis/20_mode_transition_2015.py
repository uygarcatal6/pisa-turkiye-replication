# -*- coding: utf-8 -*-
"""
20 — 2015 mod geçişi kümesinde Türkiye'nin 2012→2015 çukuru, G4 dışlama-istikrar sıralaması ve v4 şekilleri.

v4 §4.5(i), Şekil 1–2 ve Ek E G4 satırı bu betikten üretilir (önceki sürümde bu sayılar betiksizdi; etiket [E20]).

Kaynaklar (diskte; ezberden sayı yok):
  Puanlar 2012/2015: data/derived/pisa_scores_wide.csv  (köken: data/pisa/2025/annex_tables/mrq53f.xlsx, Tablo I.B1.2a.36–38)
  2015'te kâğıt tabanlı kalan sistemler: data/derived/paper_based_2015.csv (Jerrim vd. 2018, dipnot 1)
  Okul düzeyi dışlama 2018/2022: data/pisa/2018/EDU-2019-4228-EN-T010.XLSX Table I.A2.1 sütun E;
                                  data/pisa/2022/hpg9nd.xlsx Table I.A2.1 sütun E
  Şekil 1 serisi: data/derived/exclusion_frame/turkiye_exclusion_series.csv (betik 19)
Çıktılar: data/derived/mode_transition_2015/{cba2015_changes.csv, g4_exclusion_stability.csv, summary.json}
          analysis/figures/v4_sekil1_muhasebe_kapsama.png, v4_sekil2_2015_cukuru.png (öncekiler *_orig_2026-09-20.png olarak yedeklenir)
Yazan: Claude Fable 5.1 (kod), 2026-09-20. Denetim: bekliyor (Opus 4.8).
"""
import json
import os
import shutil
import numpy as np
import pandas as pd
import openpyxl
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__))).replace("\\", "/")
OUT = f"{ROOT}/data/derived/mode_transition_2015"
FIG = f"{ROOT}/analysis/figures"
os.makedirs(OUT, exist_ok=True)
REPL = "\ufffd"


def norm(s):
    return str(s).replace(REPL, "ü").replace("Turkey", "Türkiye").strip()


# ---------- 1. 2012→2015 değişimi, 2015'te bilgisayar tabanlı uygulamaya geçen sistemler ----------
sc = pd.read_csv(f"{ROOT}/data/derived/pisa_scores_wide.csv")
# 20 Eyl 2026 gece, Opus 4.8 denetimi (E19–E20): pisa_scores_wide OECD ortalama satırlarını da taşıyor (OECD, OECD_A23, OECD_A35);
# bunlar sistem değil agregadır ve "55 sistem" havuzuna girmişti (gerçek 52). Yalnız üç harfli ISO kodları tutulur.
n_before = sc.iso3.nunique()
sc = sc[sc.iso3.astype(str).str.fullmatch(r"[A-Z]{3}")].copy()
dropped_aggregates = sorted(set(pd.read_csv(f"{ROOT}/data/derived/pisa_scores_wide.csv").iso3.astype(str)) - set(sc.iso3))
print(f"agrega satırları dışarıda: {dropped_aggregates} ({n_before} → {sc.iso3.nunique()} birim)")
paper = pd.read_csv(f"{ROOT}/data/derived/paper_based_2015.csv")
paper_iso = set(paper.iso3)
a = sc[sc.cycle == 2012].set_index("iso3")
b = sc[sc.cycle == 2015].set_index("iso3")
both = a.index.intersection(b.index)
cba = [i for i in both if i not in paper_iso]
dropped_paper = sorted(i for i in both if i in paper_iso)
ch = pd.DataFrame({
    "country": b.loc[cba, "country_name"].map(norm),
    "d_math": b.loc[cba, "math_mean"] - a.loc[cba, "math_mean"],
    "d_read": b.loc[cba, "reading_mean"] - a.loc[cba, "reading_mean"],
    "d_sci": b.loc[cba, "science_mean"] - a.loc[cba, "science_mean"],
}).dropna()
ch.index.name = "iso3"
stats = {}
for dom in ["d_math", "d_read", "d_sci"]:
    s = ch[dom]
    t = float(s.loc["TUR"])
    stats[dom] = dict(tur=t, n=int(len(s)), median=float(s.median()), mean=float(s.mean()),
                      sd_ddof1=float(s.std(ddof=1)), sd_ddof0=float(s.std(ddof=0)),
                      z_ddof1=float((t - s.mean()) / s.std(ddof=1)), z_ddof0=float((t - s.mean()) / s.std(ddof=0)),
                      rank_most_negative=int((s < t).sum()) + 1,
                      second=(s.drop("TUR").idxmin(), float(s.drop("TUR").min())))
    ch[f"rank_{dom}"] = s.rank(method="min").astype(int)
ch.sort_values("d_read").to_csv(f"{OUT}/cba2015_changes.csv")

# ---------- 2. G4 — okul düzeyi dışlama sayısının 2018→2022 istikrarı ----------
def a2_school_excl(fn, sheet):
    rows = list(openpyxl.load_workbook(fn, read_only=True, data_only=True)[sheet].iter_rows(values_only=True))
    return {norm(r[0]): float(r[4]) for r in rows
            if isinstance(r[0], str) and isinstance(r[1], (int, float)) and isinstance(r[4], (int, float))}

x18 = a2_school_excl(f"{ROOT}/data/pisa/2018/EDU-2019-4228-EN-T010.XLSX", "Table I.A2.1")
x22 = a2_school_excl(f"{ROOT}/data/pisa/2022/hpg9nd.xlsx", "Table I.A2.1")
g4 = pd.DataFrame([(k, x18[k], x22[k]) for k in x18 if k in x22 and x18[k] > 0 and x22[k] > 0],
                  columns=["country", "sch_excl_2018", "sch_excl_2022"])
g4["abs_diff"] = (g4.sch_excl_2022 - g4.sch_excl_2018).abs()
g4["rel_diff_pct"] = g4.abs_diff / ((g4.sch_excl_2018 + g4.sch_excl_2022) / 2) * 100
g4 = g4.sort_values("rel_diff_pct").reset_index(drop=True)
g4["rank"] = g4.index + 1
g4.to_csv(f"{OUT}/g4_exclusion_stability.csv", index=False)
tur_g4 = g4[g4.country == "Türkiye"].iloc[0]
g4_stats = dict(n=int(len(g4)), tur_abs_diff=float(tur_g4.abs_diff), tur_rel_pct=float(tur_g4.rel_diff_pct),
                tur_rank=int(tur_g4["rank"]), second=(g4.iloc[1].country, float(g4.iloc[1].rel_diff_pct)),
                median_rel_pct=float(g4.rel_diff_pct.median()),
                enrolled_growth_pct=None)
enr18 = list(openpyxl.load_workbook(f"{ROOT}/data/pisa/2018/EDU-2019-4228-EN-T010.XLSX", read_only=True, data_only=True)["Table I.A2.1"].iter_rows(values_only=True))
enr22 = list(openpyxl.load_workbook(f"{ROOT}/data/pisa/2022/hpg9nd.xlsx", read_only=True, data_only=True)["Table I.A2.1"].iter_rows(values_only=True))
e18 = [r[2] for r in enr18 if isinstance(r[0], str) and "Turkey" in r[0]][0]
e22 = [r[2] for r in enr22 if isinstance(r[0], str) and "rkiye" in r[0]][0]
g4_stats["enrolled_growth_pct"] = float((e22 - e18) / e18 * 100)

# ---------- 3. Şekiller ----------
for name in ["v4_sekil1_muhasebe_kapsama.png", "v4_sekil2_2015_cukuru.png"]:
    src = f"{FIG}/{name}"
    bak = f"{FIG}/{name.replace('.png', '_orig_2026-09-20.png')}"
    if os.path.exists(src) and not os.path.exists(bak):
        shutil.copy2(src, bak)

ser = pd.read_csv(f"{ROOT}/data/derived/exclusion_frame/turkiye_exclusion_series.csv")
plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 11})

# Şekil 1
fig, ax = plt.subplots(figsize=(12, 6))
cyc = ser.cycle.tolist()
vals = ser.school_excl.tolist()
colors = ["#e6e4df", "#1a1a1a", "#1a1a1a", "#e6e4df"]
bars = ax.bar([str(c) for c in cyc], vals, color=colors, width=0.52, edgecolor="none")
for bi, v in zip(bars, vals):
    ax.text(bi.get_x() + bi.get_width() / 2, v + 800, f"{v:,.0f}".replace(",", "."), ha="center", fontsize=11, fontweight="bold")
ax.set_ylim(0, 54000)
ax.set_ylabel("Okul düzeyinde dışlanan öğrenci", color="#444")
ax.plot([0.75, 2.25], [47000, 47000], color="#999", lw=1.2)
ax.text(1.5, 48800, "çerçeveye yazıldı → dışlama sayıldı", ha="center", color="#888", fontsize=11, style="italic")
ax.spines[["top", "right"]].set_visible(False)
ax.set_title("Şekil 1 · Dışlama sayısı 7,6 kat oynuyor, kapsama oynamıyor", loc="left", fontsize=15, pad=14, family="DejaVu Serif")
ax2 = ax.twinx()
ax2.plot(range(4), ser.ci3, color="#8b2c1e", lw=2.4, marker="o", ms=8)
for i, v in enumerate(ser.ci3):
    ax2.text(i - 0.03, v + 0.004, f"{v:.3f}".replace(".", ","), color="#8b2c1e", ha="center", fontsize=11)
ax2.set_ylim(0.66, 0.78)
ax2.set_ylabel("Kapsama Endeksi 3 (test edilen 15 yaşlı payı)", color="#8b2c1e")
ax2.tick_params(axis="y", colors="#8b2c1e")
ax2.spines[["top"]].set_visible(False)
fig.text(0.01, 0.012, "Koyu çubuk = kurumlar çerçevede listelenip dışlandı · Açık çubuk = çerçeveye hiç alınmadı.\n"
         "Kaynak: Annex A2 (2015/18/22) · T14.A.1 (2025 taslak) · betik: analysis/20_mode_transition_2015.py [E19][E20].",
         fontsize=9.5, color="#555", va="bottom")
fig.tight_layout(rect=(0, 0.07, 1, 1))
fig.savefig(f"{FIG}/v4_sekil1_muhasebe_kapsama.png", dpi=120)
plt.close(fig)

# Şekil 2
fig, ax = plt.subplots(figsize=(13.5, 6))
doms = [("d_math", "Matematik"), ("d_sci", "Fen"), ("d_read", "Okuma")]
rng = np.random.default_rng(20260920)
for yi, (dom, label) in enumerate(doms):
    y = 2 - yi
    s = ch[dom]
    others = s.drop("TUR")
    jitter = rng.uniform(-0.12, 0.12, size=len(others))
    ax.scatter(others.values, y + jitter, s=42, color="#b9b5b0", alpha=0.9, zorder=2)
    med = float(s.median())
    ax.plot([med, med], [y - 0.27, y + 0.27], color="#8b2c1e", lw=2.2, zorder=3)
    ax.text(med, y + 0.33, f"medyan {med:.1f}".replace(".", ","), color="#8b2c1e", ha="center", fontsize=11)
    t = float(s.loc["TUR"])
    ax.scatter([t], [y], s=260, color="#111", zorder=4)
    ax.text(t, y - 0.32, f"Türkiye {t:.1f}".replace(".", ","), ha="center", fontsize=12, fontweight="bold")
ax.set_yticks([2, 1, 0])
ax.set_yticklabels([d[1] for d in doms], fontsize=13, color="#444")
ax.set_xlabel("2012 → 2015 puan değişimi", fontsize=12, color="#444")
ax.grid(axis="x", color="#ddd")
ax.spines[["top", "right"]].set_visible(False)
n = stats["d_read"]["n"]
ax.set_title(f"Şekil 2 · Aynı mod geçişini yapan {n} sistem içinde Türkiye'nin 2015 çukuru", loc="left", fontsize=15, pad=14, family="DejaVu Serif")
zs = " / ".join(f"{stats[d]['z_ddof1']:.2f}".replace(".", ",").replace("-", "−") for d in ["d_read", "d_sci", "d_math"])
rk = f"okuma {stats['d_read']['rank_most_negative']}., fen {stats['d_sci']['rank_most_negative']}., matematik {stats['d_math']['rank_most_negative']}."
fig.text(0.01, 0.012, f"Yalnız 2015'te bilgisayar tabanlı uygulamaya geçen sistemler (n={n}; kâğıt kalanlar hariç). "
         f"Türkiye en büyük düşüşte {rk} sırada; z (okuma / fen / mat, ddof=1) = {zs}.\n"
         "Kaynak: data/derived/pisa_scores_wide.csv; mod sınıflaması Jerrim vd. (2018) dipnot 1 · betik: analysis/20_mode_transition_2015.py [E20].",
         fontsize=9.5, color="#555", va="bottom")
fig.tight_layout(rect=(0, 0.08, 1, 1))
fig.savefig(f"{FIG}/v4_sekil2_2015_cukuru.png", dpi=120)
plt.close(fig)

# ---------- 4. Özet ----------
summary = dict(cba2015=dict(n=int(len(ch)), n_with_both_cycles=int(len(both)), dropped_paper_based=dropped_paper, dropped_aggregates=dropped_aggregates, stats=stats),
               g4=g4_stats,
               notes=["z: (Türkiye − küme ortalaması) / küme SD; Türkiye kümeye dâhil; ddof=1 ve ddof=0 ikisi de raporlanır.",
                      "G4 referans kümesi: 2018 ve 2022 Annex I.A2.1'de okul düzeyi dışlama sayısı sıfır olmayan sistemler; ölçüt |Δ| / iki döngü ortalaması."])
with open(f"{OUT}/summary.json", "w", encoding="utf-8") as fh:
    json.dump(summary, fh, ensure_ascii=False, indent=2, default=float)
print(json.dumps(summary, ensure_ascii=False, indent=1, default=float))
