# -*- coding: utf-8 -*-
"""
54_diag_2003_science.py — E5'in 2003 fen PV'lerinin dış doğrulaması (yayımlanmış OECD 2003 fen ülke ortalamaları ve SE'leri).

Soru (PR incelemesi, 2026-09-26): E5'in 2003 öğrenci modeli (betik 44) çekirdeğe PV1–5SCIE'yi de koyuyor, ama 2003 fen kapısı
panelde değer olmadığı için hiç sınanmadı (n = 0). Makaleye girecek [109]/[111] sayısı (2003 artığı −25,5) bu modelden geliyor.
Mikroveriden hesaplanan 2003 fen ülke ortalamaları ve SE'leri, OECD'nin yayımladığı değerlerle uyuşuyor mu?

ÖNCEDEN BELİRLENEN SEÇİMLER (RESULTS_EMPIRICAL_FIXES.md §0.13; betik koşulmadan ve yayımlanmış tablo indirilmeden önce):
  * Mikroveri: `_pisa_puf.read_columns(2003, …)`; ülke ortalaması ve SE `_pisa_puf.brr_pv_stat` (PV1–5SCIE, W_FSTUWT,
    W_FSTR1–80, Fay 0,5, Rubin) — E5'in 2003 matematik/okuma kapısıyla aynı yol.
  * Yayımlanmış değer: yazarın indirdiği resmî OECD tablosu (`extract` alt komutu; URL ve tablo adı kaydedilir, elle
    aktarım yok). Tolerans yayın hassasiyetinden: d = tablodaki en çok ondalık basamak (6'da kesilir);
    tol = max(0,001; 0,5 × 10^(−d)). Ortalama ve SE için ayrı ayrı.
  * Kapsam: E5'in 2003 sistem kümesi (`placebo_conditional_2003ps_puf.csv`, core_def mat/oku/fen, all, linear; 36).
  * Karar: "doğrulandı" ⇔ Türkiye tabloda ve PUF'ta var ve iki ölçüde de tolerans içinde, VE 36 sistemin en az 30'u eşleşti,
    VE eşleşen her sistem iki ölçüde de tolerans içinde. Aksi hâlde "doğrulanamadı" (bozan sistemler adıyla).
  * Yalnız rapor (karara girmez): ülke içi ağırlıklı korelasyonlar PV1SCIE–PV1MATH, PV1SCIE–PV1READ, PV1MATH–PV1READ
    (medyan ve Türkiye).
  * Sonuç: "doğrulandı" → E5'in 2003 fen girdisi kapanır. "doğrulanamadı" → E5'in 2003'e dayanan sayıları (−25,5 ve 2003
    karşılaştırmaları) askıya alınır, [109]/[111]'in 2003 kısmı makaleye aktarılmaz; yazara sorulur. 2025 kısmı etkilenmez.
Girdi: 2003 PUF (`_pisa_puf.REGISTRY[2003]`); indirilen OECD tablosu (.xlsx; git dışı); data/derived/analysis_panel.csv
       (ülke adı → iso3); data/derived/mechanisms/placebo_conditional_2003ps_puf.csv.
Çıktı: extract → data/derived/ontrend/pisa2003_science_published.csv, pisa2003_science_published_source.json;
       validate → data/derived/empirical_fixes/e5_science_2003_validation.csv, e5_science_2003_validation.json.
Kullanım (Windows, Git Bash; openpyxl + pyarrow):
  python analysis/54_diag_2003_science.py extract --xlsx <dosya> --sheet <sayfa> --name-col A --mean-col C --se-col D \
         --rows 10-70 --url <URL> --table "<tablo adı>" [--alias "Ülke adı=ISO3" ...]
  python analysis/54_diag_2003_science.py validate
Yazan: Claude Code (bulut oturumu), 2026-09-26. Bulutta mikroveri yok: betik sentetik veriyle sınandı.
"""
from __future__ import annotations

import argparse
import datetime as dt
import hashlib
import json
import os
import re
import sys
import unicodedata
from pathlib import Path

import numpy as np
import pandas as pd

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
sys.path.insert(0, str(HERE))
D = ROOT / "data" / "derived"
OUT = D / "empirical_fixes"
PUBDIR = D / "ontrend"
PUB_CSV = PUBDIR / "pisa2003_science_published.csv"
PUB_SRC = PUBDIR / "pisa2003_science_published_source.json"
M = 5
REPW = [f"W_FSTR{i}" for i in range(1, 81)]
MIN_MATCHED = 30
# Panelin adlarından farklı yazılan yaygın OECD 2003 adları (yalnız eşleme; eksik olan --alias ile eklenir)
ALIASES = {"Turkey": "TUR", "Czech Republic": "CZE", "Hong Kong-China": "HKG", "Macao-China": "MAC",
           "Russian Federation": "RUS", "Korea, Republic of": "KOR", "Slovakia": "SVK"}


def norm(name: str) -> str:
    s = unicodedata.normalize("NFKC", str(name)).replace(" ", " ")
    s = re.sub(r"[\d¹²³⁴⁵⁶⁷⁸⁹⁰*†‡]+$", "", s.strip())  # sondaki dipnot işaretleri
    return re.sub(r"\s+", " ", s).strip()


def parse_num(v):
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return float(v)
    if isinstance(v, str):
        s = v.strip().strip("()").replace(",", ".").strip()
        if re.fullmatch(r"-?\d+(\.\d+)?", s):
            return float(s)
    return None


def ndec(x: float) -> int:
    for d in range(7):
        if abs(x * 10 ** d - round(x * 10 ** d)) < 1e-6 * 10 ** d:
            return d
    return 6


def tol_for(values) -> tuple[int, float]:
    d = max(ndec(float(v)) for v in values)
    return d, max(0.001, 0.5 * 10 ** (-d)) + 1e-9


def name_map(extra: list[str]) -> dict:
    p = pd.read_csv(D / "analysis_panel.csv")
    m = {norm(n): i for n, i in zip(p.country_name, p.iso3) if isinstance(n, str) and len(str(i)) == 3}
    m.update({norm(k): v for k, v in ALIASES.items()})
    for a in extra:
        k, v = a.split("=", 1)
        m[norm(k)] = v.strip().upper()
    return m


def col_idx(letter: str) -> int:
    n = 0
    for ch in letter.strip().upper():
        n = n * 26 + (ord(ch) - 64)
    return n - 1


def cmd_extract(a) -> int:
    import openpyxl
    xl = Path(a.xlsx)
    wb = openpyxl.load_workbook(xl, read_only=True, data_only=True)
    ws = wb[a.sheet]
    lo, hi = (int(x) for x in a.rows.split("-")) if a.rows else (1, ws.max_row)
    nm = name_map(a.alias or [])
    ci, cm, cs = col_idx(a.name_col), col_idx(a.mean_col), col_idx(a.se_col)
    rows, unmapped = [], []
    for r_i, row in enumerate(ws.iter_rows(min_row=lo, max_row=hi, values_only=True), start=lo):
        if row is None or len(row) <= max(ci, cm, cs):
            continue
        name, mean, se = row[ci], parse_num(row[cm]), parse_num(row[cs])
        if not isinstance(name, str) or mean is None or se is None:
            continue
        iso = nm.get(norm(name))
        if iso is None:
            unmapped.append({"row": r_i, "name": name, "mean": mean})
            continue
        rows.append({"iso3": iso, "name_in_source": name, "mean": mean, "se": se, "row": r_i})
    t = pd.DataFrame(rows)
    if t.empty:
        raise SystemExit("DUR: seçilen aralıkta eşlenen satır yok (sayfa/sütun/satır aralığını kontrol et)")
    dup = t[t.iso3.duplicated(keep=False)]
    if len(dup):
        raise SystemExit(f"DUR: aynı iso3 birden çok satırda: {dup.to_dict('records')}")
    PUBDIR.mkdir(parents=True, exist_ok=True)
    tmp = PUB_CSV.with_suffix(".csv.tmp")
    t.to_csv(tmp, index=False, encoding="utf-8")
    os.replace(tmp, PUB_CSV)
    src = {"xlsx_file": xl.name, "xlsx_sha256": hashlib.sha256(xl.read_bytes()).hexdigest(), "sheet": a.sheet,
           "name_col": a.name_col, "mean_col": a.mean_col, "se_col": a.se_col, "rows": a.rows, "url": a.url, "table": a.table,
           "aliases_cli": a.alias or [], "n_rows": int(len(t)), "unmapped_numeric_rows": unmapped,
           "extracted_at": dt.datetime.now().isoformat(timespec="seconds")}
    tmp = PUB_SRC.with_suffix(".json.tmp")
    tmp.write_text(json.dumps(src, ensure_ascii=False, indent=1, allow_nan=False), encoding="utf-8")
    os.replace(tmp, PUB_SRC)
    print(f"{len(t)} satır yazıldı → {PUB_CSV}")
    print(f"eşlenemeyen sayısal satır: {len(unmapped)}" + (f" (ilk 10: {[u['name'] for u in unmapped[:10]]})" if unmapped else ""))
    return 0


def load_2003() -> pd.DataFrame:
    import _pisa_puf as P
    cols = ["CNT", "W_FSTUWT"] + REPW + [f"PV{i}{r}" for r in ("SCIE", "MATH", "READ") for i in range(1, M + 1)]
    df = P.read_columns(2003, cols)
    df["CNT"] = df["CNT"].astype(str).str.strip()
    for c in df.columns:
        if c != "CNT":
            df[c] = pd.to_numeric(df[c], errors="coerce")
    return df


def wcorr(x, y, w) -> float:
    ok = np.isfinite(x) & np.isfinite(y) & np.isfinite(w)
    x, y, w = x[ok], y[ok], w[ok]
    mx, my = np.average(x, weights=w), np.average(y, weights=w)
    return float(np.sum(w * (x - mx) * (y - my)) / np.sqrt(np.sum(w * (x - mx) ** 2) * np.sum(w * (y - my) ** 2)))


def cmd_validate(a) -> int:
    import _pisa_puf as P
    if not PUB_CSV.exists():
        raise SystemExit(f"DUR: {PUB_CSV} yok; önce `extract` koş")
    pub = pd.read_csv(PUB_CSV).set_index("iso3")
    src = json.loads(PUB_SRC.read_text(encoding="utf-8")) if PUB_SRC.exists() else {}
    d_m, tol_m = tol_for(pub["mean"])
    d_s, tol_s = tol_for(pub["se"])
    s = pd.read_csv(D / "mechanisms" / "placebo_conditional_2003ps_puf.csv", encoding="utf-8-sig")
    e5 = sorted(s[(s.core_def == "mat/oku/fen") & (s.scope == "all") & (s.spec == "linear")].iso3.unique())
    df = load_2003()
    est = P.brr_pv_stat(df, [f"PV{i}SCIE" for i in range(1, M + 1)], "W_FSTUWT", REPW).set_index("iso3")
    rows = []
    for iso in sorted(set(e5) | (set(pub.index) & set(est.index))):
        r = {"iso3": iso, "in_e5": iso in e5, "in_published": iso in pub.index, "in_puf": iso in est.index}
        if r["in_published"] and r["in_puf"]:
            dm = float(est.loc[iso, "mean"] - pub.loc[iso, "mean"])
            dse = float(est.loc[iso, "se"] - pub.loc[iso, "se"])
            r.update({"mean_micro": round(float(est.loc[iso, "mean"]), 4), "mean_pub": float(pub.loc[iso, "mean"]),
                      "diff_mean": round(dm, 4), "se_micro": round(float(est.loc[iso, "se"]), 4), "se_pub": float(pub.loc[iso, "se"]),
                      "diff_se": round(dse, 4), "ok_mean": bool(abs(dm) <= tol_m), "ok_se": bool(abs(dse) <= tol_s)})
        rows.append(r)
    t = pd.DataFrame(rows)
    m5 = t[t.in_e5 & t.in_published & t.in_puf]
    ok = m5.ok_mean.astype(bool) & m5.ok_se.astype(bool)  # sütun object tipinde olabilir; ~ bool üzerinde çalışsın
    bad = m5[~ok]
    tur = t[t.iso3 == "TUR"]
    tur_ok = bool(len(tur) and tur.iloc[0].get("in_published") and tur.iloc[0].get("in_puf")
                  and bool(tur.iloc[0].ok_mean) and bool(tur.iloc[0].ok_se))
    passed = tur_ok and len(m5) >= MIN_MATCHED and len(bad) == 0
    cors = []
    for iso in e5:
        g = df[df.CNT == iso]
        if len(g) < 2:
            continue
        w = g.W_FSTUWT.to_numpy(float)
        v = {k: g[f"PV1{k}"].to_numpy(float) for k in ("SCIE", "MATH", "READ")}
        cors.append({"iso3": iso, "sci_math": wcorr(v["SCIE"], v["MATH"], w), "sci_read": wcorr(v["SCIE"], v["READ"], w),
                     "math_read": wcorr(v["MATH"], v["READ"], w)})
    c = pd.DataFrame(cors).set_index("iso3") if cors else pd.DataFrame()
    info = {"note": "E5 2003 fen PV'lerinin dış doğrulaması; ön-belirleme §0.13.", "source": src,
            "decimals": {"mean": d_m, "se": d_s}, "tolerance": {"mean": round(tol_m, 6), "se": round(tol_s, 6)},
            "e5_systems": len(e5), "e5_matched": int(len(m5)), "min_matched": MIN_MATCHED,
            "e5_missing_in_published": sorted(t[t.in_e5 & ~t.in_published].iso3), "e5_missing_in_puf": sorted(t[t.in_e5 & ~t.in_puf].iso3),
            "failing": bad[["iso3", "diff_mean", "diff_se"]].to_dict("records"), "tur_ok": tur_ok,
            "max_abs_diff_mean_e5": (round(float(m5.diff_mean.abs().max()), 4) if len(m5) else None),
            "max_abs_diff_se_e5": (round(float(m5.diff_se.abs().max()), 4) if len(m5) else None),
            "correlations_report_only": ({k: {"median": round(float(c[k].median()), 4),
                                              "TUR": (round(float(c.loc["TUR", k]), 4) if "TUR" in c.index else None)}
                                          for k in c.columns} if len(c) else {}),
            "verdict": ("doğrulandı: 2003 fen PV'leri yayımlanmış değerleri yayın hassasiyetinde üretiyor" if passed else
                        "doğrulanamadı: E5'in 2003'e dayanan sayıları askıya alınır; yazara sorulur")}
    OUT.mkdir(parents=True, exist_ok=True)
    tmp = OUT / "e5_science_2003_validation.csv.tmp"
    t.to_csv(tmp, index=False, encoding="utf-8")
    os.replace(tmp, OUT / "e5_science_2003_validation.csv")
    tmp = OUT / "e5_science_2003_validation.json.tmp"
    tmp.write_text(json.dumps(info, ensure_ascii=False, indent=1, allow_nan=False, default=float), encoding="utf-8")
    os.replace(tmp, OUT / "e5_science_2003_validation.json")
    print(json.dumps({k: info[k] for k in ("decimals", "tolerance", "e5_matched", "failing", "tur_ok", "verdict")},
                     ensure_ascii=False, indent=1, default=float))
    return 0 if passed else 2


def main() -> int:
    ap = argparse.ArgumentParser()
    sub = ap.add_subparsers(dest="cmd", required=True)
    e = sub.add_parser("extract")
    e.add_argument("--xlsx", required=True)
    e.add_argument("--sheet", required=True)
    e.add_argument("--name-col", required=True)
    e.add_argument("--mean-col", required=True)
    e.add_argument("--se-col", required=True)
    e.add_argument("--rows", default=None, help="ör. 10-70")
    e.add_argument("--url", required=True)
    e.add_argument("--table", required=True)
    e.add_argument("--alias", action="append", help='"Ülke adı=ISO3" (birden çok kez)')
    sub.add_parser("validate")
    a = ap.parse_args()
    return cmd_extract(a) if a.cmd == "extract" else cmd_validate(a)


if __name__ == "__main__":
    sys.exit(main())
